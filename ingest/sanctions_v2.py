"""Labels v2 (task P4-9): adds vessel designations from sources v1 (``ingest/sanctions.py``:
OFAC/EU-FSF/UK) misses, without touching v1 at all -- v1 stays the frozen pre-registered snapshot
(``data/reference/sanctions.parquet``, real run 2026-09-21). This module lands a second table,
``data/reference/sanctions_v2.parquet``, that is every v1 row verbatim plus new rows from three
additional sources, all normalised to the same :class:`ingest.sanctions.SanctionedVessel` shape.

**Sources landed, all confirmed live 2026-09-29:**

1. **EU vessel designations** (``fetch_eu_vessels``, ``source="eu_vessels"``) -- the shadow-fleet
   port-access/services ban under Article 3s of Council Regulation (EU) 833/2014 (the measure
   Annex XLII lists vessels under). Sourced from the **Danish Maritime Authority**'s own published
   list, not OpenSanctions: :data:`EU_VESSELS_PAGE_URL` is fetched first and its ``.csv`` link
   extracted by regex, rather than hardcoding the link target directly, because that target
   embeds a media ID (e.g. ``.../639204766104358805/...``) that changes on every DMA update --
   confirmed by the filename's own date-like suffix (``...240726.csv``, i.e. 2026-07-24, matching
   this snapshot's newest designation date). The CSV is semicolon-delimited, UTF-8 with a BOM, and
   its IMO column header carries a trailing space (``"IMO number "``, confirmed live) -- a real
   quirk, not a typo here. Rows are filtered to ``Subject to`` starting with ``"Article 3s"``
   (Council Regulation 833/2014); the same file also carries unrelated DPRK ship designations
   under Council Regulation 2017/1509 (``Asset freeze``/``Seizure``/``Port entry ban``/
   ``De-registration``), which are out of this task's scope and dropped. **Cross-checked 2026-09-29
   against OpenSanctions' ``eu_sanctions`` dataset** (not ``eu_fsf`` -- the port-access ban is not
   a financial-sanctions asset freeze, so it lives in OpenSanctions' broader aggregate dataset, not
   the FSF one v1 already uses; confirmed by grep: 0 ``EU-MARE``-tagged rows in ``eu_fsf``, 673 in
   ``eu_sanctions``): 671 distinct IMOs in common between DMA and OpenSanctions'
   ``EU-MARE``-tagged, IMO-bearing vessel rows, **671/671 (100%) exact designation-date agreement**;
   DMA's IMO set is a strict subset of OpenSanctions' (OpenSanctions carries 2 extra IMOs, of which
   1 has no parseable date, so it wouldn't have been landed here either way). This is strong
   corroboration and DMA is used as the single source of record (landing both would double-count
   the same designations) -- the cross-check itself is not re-run by this module, it was a
   one-off verification (see the P4-9 session's report).

2. **Canada SEMA** (``fetch_ca``, ``source="ca"``) -- Global Affairs Canada's consolidated
   Special Economic Measures Act list, the official XML export
   (:data:`CA_SEMA_URL`, ~2.9MB, confirmed live). A flat ``<record>`` list (no nested containers
   needing multi-pass resolution the way OFAC's XML does), so this is parsed with a plain
   ``xml.etree.ElementTree.parse`` rather than ``iterparse`` -- deliberately simpler than
   :mod:`ingest.sanctions`'s OFAC parser, not an oversight; see that module's docstring for when
   streaming is warranted. A ship record is one with a non-empty
   ``ShipIMONumber-NumeroOMIDuNavire`` element; the sibling ``DateOfListing-DateDinscription`` is
   already ISO ``YYYY-MM-DD``. **Real quirk found while verifying, not in any spec:** 5 records
   under Schedule 1 Part 2 (Iran individuals, confirmed by ``Country-Pays``) carry Persian-script
   text in the ``ShipIMONumber`` field instead of a real IMO -- a data-entry/export bug on GAC's
   side, not this parser's -- caught by the same checksum gate (below) as any other malformed IMO,
   not specially handled.

3. **New Zealand** MFAT Russia Sanctions Register (``fetch_nz``, ``source="nz"``) -- the register's
   own ``Ships`` worksheet (:data:`NZ_REGISTER_URL`, ~21MB .xlsx, confirmed live), which has a
   real per-vessel ``Date of Sanction`` column plus ``Unique Identifier`` (used as ``source_id``
   directly -- a real stable key, unlike the other two new sources which have none of their own).
   Read with ``openpyxl`` in ``read_only=True`` mode (lazy row iteration, never the whole workbook
   materialised) rather than loading via pandas, since the file carries four other sheets
   (``Russia Sanctions Register``, ``Trade Measures - Oil Price Cap``, ``Trade Measures``,
   ``Exemptions``) this module has no use for and never touches. A row with
   ``Record Deleted Flag`` set is excluded and counted separately (``deleted``, not the same
   bucket as ``no_date``/``invalid_imo``/``after_cap`` below) -- not observed on the live 2026-09-29
   snapshot (0/210), but real per the sheet's own schema, so handled defensively and covered by a
   synthetic test rather than left untested per this project's convention (an untested branch is as
   real a gap as a wrong one).

**Australia (DFAT) and Switzerland (SECO): not landed.** DFAT's site
(``www.dfat.gov.au``, both the homepage and the consolidated-list page) silently drops the TCP
connection after the TLS handshake completes -- confirmed 2026-09-29 from this project's network,
multiple attempts, both ``httpx`` and a raw ``curl -v`` (which shows the request sent and then
nothing, not even a timeout at the HTTP layer) -- consistent with a bot-management challenge that
never resolves for non-browser clients, not a transient outage (static asset paths under
``/sites/default/files/`` answer instantly with a real 404). No workaround was in scope. SECO's
consolidated sanctions data (``sesam.search.admin.ch``) is a search UI, not a bulk export; no
machine-readable per-vessel-with-real-date file was found within this session's search budget, and
the task itself treats Switzerland as lowest-priority ("mostly mirrors the EU"). Both are left for
a future session if pursued further; see the P4-9 session's report for the exact URLs tried.

**Exclusion gates, applied to every new-source record, matching v1's IMO checksum convention
(:data:`process.identity.VALID_IMO_SQL`, reused verbatim via DuckDB rather than reimplemented in
Python -- same precedent as ``process/sanctions_match.py``) and this task's own date rules:**

1. ``no_date`` -- no parseable per-record designation date. Unlike v1's OFAC/UK parsers (which
   raise on a missing date, since v1 treats that as a schema violation on files where every row is
   expected to carry one), v2's new sources treat a missing date as an ordinary, expected
   exclusion reason: count it and move on, since these are less-audited feeds than the big-three.
2. ``invalid_imo`` -- fails :data:`process.identity.VALID_IMO_SQL`'s format/checksum (this also
   catches an empty/missing IMO, which trivially fails the same check).
3. ``after_cap`` -- designation date strictly after :data:`V1_SNAPSHOT_CUTOFF` (2026-09-21, v1's
   own real fetch date) -- keeps v2 a fair extension of v1's forward-looking window, not a source
   of look-ahead a future evaluation could accidentally use.

A record can be excluded for at most one reason (checked in the order above); each source function
returns ``(list[SanctionedVessel], dict[str, int])`` -- the vessels kept and a count per exclusion
reason plus ``kept`` -- rather than only a list, since counting is part of this task's spec, not an
afterthought logged and discarded.

**Landing** (:func:`build_sanctions_v2`). Reads v1 (``data/reference/sanctions.parquet``) with
DuckDB and re-lands every column unchanged (including each row's own original ``built_at``) --
this module never re-fetches or re-derives a v1 row. New-source rows get one shared ``built_at``
timestamp for this v2 build run. Atomic write via
:func:`process.partitions.atomic_write_parquet` (temp file + ``os.replace``), same as every other
Parquet-producing module in this project. Idempotent like ``ingest.sanctions.build_sanctions``: a
no-op if the output already exists, unless ``force=True``.
"""

from __future__ import annotations

import argparse
import csv
import logging
import re
import tempfile
import xml.etree.ElementTree as ET
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path

import duckdb
import httpx
import openpyxl

from ingest.sanctions import (
    SANCTIONS_PATH,
    SanctionedVessel,
    SanctionsSchemaError,
    _stream_download,
)
from process.identity import VALID_IMO_SQL
from process.partitions import atomic_write_parquet

logger = logging.getLogger(__name__)

SANCTIONS_V2_PATH = SANCTIONS_PATH.with_name("sanctions_v2.parquet")

# v1's own real fetch date (ingest.sanctions module docstring: "Real run 2026-09-21"). A
# new-source record designated after this is excluded -- see module docstring's exclusion gates.
V1_SNAPSHOT_CUTOFF = date(2026, 9, 21)

EU_VESSELS_PAGE_URL = (
    "https://www.dma.dk/growth-and-framework-conditions/maritime-sanctions/"
    "general-information/eu-vessel-designations"
)
CA_SEMA_URL = (
    "https://www.international.gc.ca/world-monde/assets/office_docs/"
    "international_relations-relations_internationales/sanctions/sema-lmes.xml"
)
NZ_REGISTER_URL = (
    "https://www.mfat.govt.nz/assets/Countries-and-Regions/Europe/Ukraine/"
    "Russia-Sanctions-Register.xlsx"
)

_EU_VESSELS_CSV_LINK_RE = re.compile(r'href="([^"]+\.csv)"', re.IGNORECASE)
# The DMA CSV's own header, trailing space on the IMO column confirmed live -- see module
# docstring.
_EU_VESSELS_REQUIRED_COLUMNS = {"Vessel name", "IMO number ", "Date of application", "Subject to"}
_EU_VESSELS_PROGRAM_PREFIX = "Article 3s"  # Council Regulation (EU) 833/2014's port-access ban


class EUVesselsSchemaError(SanctionsSchemaError):
    """The DMA EU-vessel-designations page/CSV does not match what was confirmed live 2026-09-29."""


class CASchemaError(SanctionsSchemaError):
    """The Canada SEMA XML structure does not match what was confirmed live 2026-09-29."""


class NZSchemaError(SanctionsSchemaError):
    """The NZ MFAT Russia Sanctions Register's 'Ships' sheet does not match what was confirmed
    live 2026-09-29."""


@dataclass(frozen=True)
class _Candidate:
    """One not-yet-validated designation, before the shared exclusion gates run -- see module
    docstring. ``imo``/``designation_date`` are deliberately permissive (``imo`` may be empty or
    garbage, ``designation_date`` may be ``None``): validating them is the gates' job, not the
    per-source parser's.
    """

    source_id: str
    vessel_name: str
    imo: str
    flag: str | None
    program: str
    designation_date: date | None


def _valid_imo_set(imos: set[str]) -> set[str]:
    """Return the subset of imos (bare strings, possibly empty or non-digit) that pass this
    project's IMO checksum -- :data:`process.identity.VALID_IMO_SQL`, reused verbatim rather than
    reimplemented in Python, same precedent as ``process/sanctions_match.py``.
    """
    if not imos:
        return set()
    con = duckdb.connect()
    try:
        con.execute("CREATE TEMP TABLE _candidate_imos (imo VARCHAR)")
        con.executemany("INSERT INTO _candidate_imos VALUES (?)", [(i,) for i in imos])
        rows = con.execute(f"SELECT imo FROM _candidate_imos WHERE {VALID_IMO_SQL}").fetchall()
        return {r[0] for r in rows}
    finally:
        con.close()


def _apply_gates(
    candidates: list[_Candidate], source: str
) -> tuple[list[SanctionedVessel], dict[str, int]]:
    """Run every candidate through the three shared exclusion gates (module docstring), in order,
    and normalise the survivors to :class:`SanctionedVessel`.
    """
    counts = {"total": len(candidates), "no_date": 0, "invalid_imo": 0, "after_cap": 0, "kept": 0}

    dated = [c for c in candidates if c.designation_date is not None]
    counts["no_date"] = len(candidates) - len(dated)

    valid_imos = _valid_imo_set({c.imo for c in dated})

    result: list[SanctionedVessel] = []
    for c in dated:
        if c.imo not in valid_imos:
            counts["invalid_imo"] += 1
            continue
        if c.designation_date > V1_SNAPSHOT_CUTOFF:
            counts["after_cap"] += 1
            continue
        counts["kept"] += 1
        result.append(
            SanctionedVessel(
                source=source,
                source_id=c.source_id,
                vessel_name=c.vessel_name,
                imo=c.imo,
                flag=c.flag,
                program=c.program,
                designation_date=c.designation_date,
                designation_date_precision="day",
            )
        )
    logger.info(
        "%s: %d total, %d kept, %d excluded (no_date=%d invalid_imo=%d after_cap=%d)",
        source,
        counts["total"],
        counts["kept"],
        counts["total"] - counts["kept"],
        counts["no_date"],
        counts["invalid_imo"],
        counts["after_cap"],
    )
    return result, counts


# ---------------------------------------------------------------------------
# EU vessel designations (DMA)
# ---------------------------------------------------------------------------


def _discover_eu_vessels_csv_url(page_text: str, page_url: str) -> str:
    """Extract the '.csv' download link from the DMA page's own HTML -- see module docstring for
    why this is discovered dynamically rather than hardcoded (the link embeds a media ID that
    changes on every DMA update).
    """
    match = _EU_VESSELS_CSV_LINK_RE.search(page_text)
    if match is None:
        raise EUVesselsSchemaError(
            f"No '.csv' link found on {page_url!r} -- the DMA page structure has changed since "
            "this module was last verified (2026-09-29), see module docstring."
        )
    href = match.group(1)
    if href.startswith("/"):
        return "https://www.dma.dk" + href
    return href


def _parse_eu_vessels_csv(path: Path) -> tuple[list[SanctionedVessel], dict[str, int]]:
    """Parse the DMA EU-vessel-designations CSV into normalised :class:`SanctionedVessel`
    records. See module docstring for the column-by-column mapping, the trailing-space IMO
    column quirk, and the ``Subject to`` filter that keeps only Article 3s (Council Regulation
    833/2014) rows -- the file also carries unrelated DPRK ship designations.
    """
    with open(path, encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        fieldnames = set(reader.fieldnames or [])
        if not _EU_VESSELS_REQUIRED_COLUMNS.issubset(fieldnames):
            raise EUVesselsSchemaError(
                f"Expected columns {sorted(_EU_VESSELS_REQUIRED_COLUMNS)} in the DMA EU vessel "
                f"CSV, found {sorted(fieldnames)} -- see module docstring."
            )
        rows = [
            row
            for row in reader
            if (row.get("Subject to") or "").strip().startswith(_EU_VESSELS_PROGRAM_PREFIX)
        ]

    # A vessel can appear more than once (re-listed under a later package with its own date, 3 of
    # 674 real rows did on 2026-09-29) -- earliest wins, same convention as v1's OFAC/UK parsers.
    rows_by_imo: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        imo = (row.get("IMO number ") or "").strip().removeprefix("IMO").strip()
        rows_by_imo[imo].append(row)

    candidates: list[_Candidate] = []
    for imo, group in rows_by_imo.items():
        dated_rows: list[tuple[date, dict]] = []
        for row in group:
            raw = (row.get("Date of application") or "").strip()
            if not raw:
                continue
            try:
                d = datetime.strptime(raw, "%d.%m.%Y").date()  # noqa: DTZ007 -- date-only field
            except ValueError:
                continue
            dated_rows.append((d, row))
        if dated_rows:
            d, earliest_row = min(dated_rows, key=lambda pair: pair[0])
        else:
            d, earliest_row = None, group[0]
        name = (earliest_row.get("Vessel name") or "").strip()
        program = (earliest_row.get("Subject to") or "").strip() or "UNKNOWN"
        candidates.append(
            _Candidate(
                source_id=f"dma-eu-{imo or 'unknown'}",
                vessel_name=name,
                imo=imo,
                flag=None,
                program=program,
                designation_date=d,
            )
        )
    return _apply_gates(candidates, source="eu_vessels")


def fetch_eu_vessels(
    client: httpx.Client | None = None, tmp_dir: Path | None = None
) -> tuple[list[SanctionedVessel], dict[str, int]]:
    """Download and parse the DMA's EU vessel designations list. See module docstring."""
    owns_client = client is None
    client = client or httpx.Client(timeout=60.0, follow_redirects=True)
    try:
        page = client.get(EU_VESSELS_PAGE_URL)
        page.raise_for_status()
        csv_url = _discover_eu_vessels_csv_url(page.text, EU_VESSELS_PAGE_URL)
        with tempfile.TemporaryDirectory(prefix="eu-vessels-", dir=tmp_dir) as tmp:
            csv_path = Path(tmp) / "eu_vessel_designations.csv"
            _stream_download(client, csv_url, csv_path)
            vessels, counts = _parse_eu_vessels_csv(csv_path)
    finally:
        if owns_client:
            client.close()
    return vessels, counts


# ---------------------------------------------------------------------------
# Canada SEMA
# ---------------------------------------------------------------------------


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def parse_ca_sema_xml(path: Path) -> tuple[list[SanctionedVessel], dict[str, int]]:
    """Parse Global Affairs Canada's consolidated SEMA XML into normalised
    :class:`SanctionedVessel` records. See module docstring -- a plain ``ET.parse`` (the file is
    small and flat, no multi-pass reference resolution needed the way OFAC's XML does).
    """
    try:
        tree = ET.parse(path)
    except ET.ParseError as exc:
        raise CASchemaError(f"{path} is not well-formed XML -- see module docstring.") from exc
    root = tree.getroot()

    rows_by_imo: dict[str, list[ET.Element]] = defaultdict(list)
    for rec in root.iter():
        if _local(rec.tag) != "record":
            continue
        imo_elem = next(
            (e for e in rec if _local(e.tag) == "ShipIMONumber-NumeroOMIDuNavire"), None
        )
        imo_text = (imo_elem.text or "").strip() if imo_elem is not None else ""
        if not imo_text:
            continue  # not a ship record at all (person/entity) -- not a gate, a structural filter
        rows_by_imo[imo_text].append(rec)

    candidates: list[_Candidate] = []
    for imo, group in rows_by_imo.items():
        dated_rows: list[tuple[date, ET.Element]] = []
        for rec in group:
            date_elem = next(
                (e for e in rec if _local(e.tag) == "DateOfListing-DateDinscription"), None
            )
            raw = (date_elem.text or "").strip() if date_elem is not None else ""
            if not raw:
                continue
            try:
                d = date.fromisoformat(raw)
            except ValueError:
                continue
            dated_rows.append((d, rec))
        if dated_rows:
            d, earliest = min(dated_rows, key=lambda pair: pair[0])
        else:
            d, earliest = None, group[0]

        def _text(rec: ET.Element, tag: str) -> str | None:
            elem = next((e for e in rec if _local(e.tag) == tag), None)
            return (elem.text or "").strip() if elem is not None and elem.text else None

        name = _text(earliest, "EntityOrShip-EntiteOuNavire") or ""
        schedule = _text(earliest, "Schedule-Annexe")
        item = _text(earliest, "Item-NumeroDarticle")
        program = f"SEMA Schedule {schedule}" if schedule else "UNKNOWN"
        source_id = f"ca-{schedule}-{item}" if schedule and item else f"ca-{imo}"
        candidates.append(
            _Candidate(
                source_id=source_id,
                vessel_name=name,
                imo=imo,
                flag=None,  # Country-Pays is the targeted regime, not the vessel's actual flag
                program=program,
                designation_date=d,
            )
        )
    return _apply_gates(candidates, source="ca")


def fetch_ca(
    client: httpx.Client | None = None, tmp_dir: Path | None = None
) -> tuple[list[SanctionedVessel], dict[str, int]]:
    """Download and parse Canada's consolidated SEMA XML. See module docstring."""
    owns_client = client is None
    client = client or httpx.Client(timeout=120.0, follow_redirects=True)
    try:
        with tempfile.TemporaryDirectory(prefix="ca-sema-", dir=tmp_dir) as tmp:
            xml_path = Path(tmp) / "sema-lmes.xml"
            _stream_download(client, CA_SEMA_URL, xml_path)
            vessels, counts = parse_ca_sema_xml(xml_path)
    finally:
        if owns_client:
            client.close()
    return vessels, counts


# ---------------------------------------------------------------------------
# New Zealand MFAT Russia Sanctions Register
# ---------------------------------------------------------------------------

_NZ_SHIPS_SHEET = "Ships"
_NZ_REQUIRED_COLUMNS = {
    "Unique Identifier",
    "IMO Number",
    "Name of Ship as of Date of Sanction",
    "Date of Sanction",
    "Record Deleted Flag",
}


def _nz_cell_to_imo(value: object) -> str:
    """Normalise one 'IMO Number' cell (openpyxl hands back an int/float for a numeric cell, or a
    str if the workbook stored it as text) to a bare-digit string, zero-padded defensively (no
    real IMO is short enough to need it, confirmed live, but a leading-zero IMO is not impossible
    per the checksum's own math)."""
    if value is None:
        return ""
    if isinstance(value, float):
        value = int(value)
    text = str(value).strip().removeprefix("IMO").strip()
    return text.zfill(7) if text.isdigit() else text


def _nz_cell_to_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.strptime(text, "%d/%m/%Y").date()  # noqa: DTZ007 -- date-only field
    except ValueError:
        return None


def parse_nz_register_xlsx(path: Path) -> tuple[list[SanctionedVessel], dict[str, int]]:
    """Parse the NZ MFAT Russia Sanctions Register's 'Ships' sheet into normalised
    :class:`SanctionedVessel` records. See module docstring -- ``openpyxl`` read-only (lazy row
    iteration), 'Unique Identifier' used directly as ``source_id`` (a real stable key this source
    has and the other two new sources don't), deleted records excluded and counted separately.
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover -- exercised only by a genuinely corrupt file
        raise NZSchemaError(f"{path} is not a readable .xlsx file -- see module docstring.") from exc
    try:
        if _NZ_SHIPS_SHEET not in wb.sheetnames:
            raise NZSchemaError(
                f"Expected a {_NZ_SHIPS_SHEET!r} sheet in the NZ register, found "
                f"{wb.sheetnames} -- see module docstring."
            )
        ws = wb[_NZ_SHIPS_SHEET]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            raise NZSchemaError(f"{_NZ_SHIPS_SHEET!r} sheet has no header row.")
        idx = {name: i for i, name in enumerate(header) if name is not None}
        if not _NZ_REQUIRED_COLUMNS.issubset(idx):
            raise NZSchemaError(
                f"Expected columns {sorted(_NZ_REQUIRED_COLUMNS)} in the NZ register's "
                f"{_NZ_SHIPS_SHEET!r} sheet, found {sorted(idx)} -- see module docstring."
            )

        deleted = 0
        candidates: list[_Candidate] = []
        for row in rows_iter:
            if row[idx["Record Deleted Flag"]]:
                deleted += 1
                continue
            imo = _nz_cell_to_imo(row[idx["IMO Number"]])
            name = str(row[idx["Name of Ship as of Date of Sanction"]] or "").strip()
            d = _nz_cell_to_date(row[idx["Date of Sanction"]])
            source_id = str(row[idx["Unique Identifier"]] or "").strip() or f"nz-{imo}"
            candidates.append(
                _Candidate(
                    source_id=source_id,
                    vessel_name=name,
                    imo=imo,
                    flag=None,
                    program="Russia Sanctions Act 2022 (Ships)",
                    designation_date=d,
                )
            )
    finally:
        wb.close()

    vessels, counts = _apply_gates(candidates, source="nz")
    counts["deleted"] = deleted
    if deleted:
        logger.info("nz: %d additional record(s) excluded (deleted)", deleted)
    return vessels, counts


def fetch_nz(
    client: httpx.Client | None = None, tmp_dir: Path | None = None
) -> tuple[list[SanctionedVessel], dict[str, int]]:
    """Download and parse the NZ MFAT Russia Sanctions Register's 'Ships' sheet. See module
    docstring."""
    owns_client = client is None
    client = client or httpx.Client(timeout=120.0, follow_redirects=True)
    try:
        with tempfile.TemporaryDirectory(prefix="nz-register-", dir=tmp_dir) as tmp:
            xlsx_path = Path(tmp) / "russia-sanctions-register.xlsx"
            _stream_download(client, NZ_REGISTER_URL, xlsx_path)
            vessels, counts = parse_nz_register_xlsx(xlsx_path)
    finally:
        if owns_client:
            client.close()
    return vessels, counts


# ---------------------------------------------------------------------------
# Landing
# ---------------------------------------------------------------------------


def build_sanctions_v2(
    v1_path: Path = SANCTIONS_PATH,
    out_path: Path = SANCTIONS_V2_PATH,
    client: httpx.Client | None = None,
    tmp_dir: Path | None = None,
    force: bool = False,
) -> tuple[Path, dict[str, dict[str, int]]]:
    """Land v2: every v1 row unchanged, plus new rows from the three sources above. Idempotent
    like :func:`ingest.sanctions.build_sanctions`: a no-op if out_path already exists, unless
    force=True. Returns (out_path, {source: exclusion-counts}) for the three new sources -- v1's
    own rows carry no exclusion counts here, since they are read, not re-derived.
    """
    if out_path.exists() and not force:
        logger.info(
            "%s already exists, skipping (pass force=True / --force to re-fetch)", out_path
        )
        return out_path, {}
    if not v1_path.exists():
        raise FileNotFoundError(
            f"{v1_path} does not exist -- run `python -m ingest.sanctions` first; v2 extends v1, "
            "it does not build it."
        )

    all_counts: dict[str, dict[str, int]] = {}
    new_vessels: list[SanctionedVessel] = []
    for source_name, fetch in (
        ("eu_vessels", fetch_eu_vessels),
        ("ca", fetch_ca),
        ("nz", fetch_nz),
    ):
        vessels, counts = fetch(client=client, tmp_dir=tmp_dir)
        new_vessels.extend(vessels)
        all_counts[source_name] = counts
    logger.info("Total new-source vessel record(s) across all three sources: %d", len(new_vessels))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect()
    try:
        con.execute(
            "CREATE OR REPLACE TEMP TABLE _new (source VARCHAR, source_id VARCHAR, "
            "vessel_name VARCHAR, imo VARCHAR, flag VARCHAR, program VARCHAR, "
            "designation_date DATE, designation_date_precision VARCHAR)"
        )
        if new_vessels:
            con.executemany(
                "INSERT INTO _new VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                [
                    (
                        v.source,
                        v.source_id,
                        v.vessel_name,
                        v.imo,
                        v.flag,
                        v.program,
                        v.designation_date,
                        v.designation_date_precision,
                    )
                    for v in new_vessels
                ],
            )
        built_at = datetime.now(timezone.utc)
        select_sql = (
            "SELECT source, source_id, vessel_name, imo, flag, program, designation_date, "
            f"designation_date_precision, built_at FROM read_parquet('{v1_path.as_posix()}') "
            "UNION ALL "
            "SELECT source, source_id, vessel_name, imo, flag, program, designation_date, "
            f"designation_date_precision, TIMESTAMP '{built_at.strftime('%Y-%m-%d %H:%M:%S.%f')}' "
            "AS built_at FROM _new "
            "ORDER BY source, designation_date"
        )
        atomic_write_parquet(con, select_sql, out_path)
    finally:
        con.close()
    return out_path, all_counts


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build labels v2: v1's sanctions.parquet plus EU/CA/NZ new-source vessel "
        "designations, landed at data/reference/sanctions_v2.parquet."
    )
    parser.add_argument(
        "--v1-path", default=str(SANCTIONS_PATH), help="Path to the frozen v1 table"
    )
    parser.add_argument(
        "--out-path", default=str(SANCTIONS_V2_PATH), help="Output path for the v2 table"
    )
    parser.add_argument(
        "--force", action="store_true", help="Re-fetch and rebuild even if the output already exists"
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    args = _parse_args(argv)
    _, counts = build_sanctions_v2(
        v1_path=Path(args.v1_path), out_path=Path(args.out_path), force=args.force
    )
    for source, source_counts in counts.items():
        logger.info("%s: %s", source, source_counts)


if __name__ == "__main__":
    main()
